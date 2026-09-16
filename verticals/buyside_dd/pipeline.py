"""
Buyside DD Pipeline orchestrator.

Usage:
  python3 pipeline.py inputs/your_data_room/

Stages:
  1. Load documents from inputs/
  2. Extract claims (LLM)
  3. Type referents (LLM + ontology)
  4. Look up f-rules (f-library)
  5. Look up M sources (source atlas)
  6. Acquire observations (connectors)
  7. Compare and score
  8. Write divergence report to outputs/
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Optional

# Pipeline-relative imports
sys.path.insert(0, str(Path(__file__).parent.parent.parent))
from verticals.buyside_dd.schemas import (
    Claim, TypedClaim, ClaimWithF, ClaimWithSources,
    Finding, Observation, FRule, MSource, ReferentType, Severity, Verifiability,
)
from verticals.buyside_dd.f_library import find_rules, REAL_ESTATE_RULES
from verticals.buyside_dd.source_atlas import (
    find_sources, REAL_ESTATE_SOURCES, recall_floor_gaps, source_by_id,
)
from verticals.buyside_dd.dispatcher import dispatch_with_policy, coverage_report
from verticals.buyside_dd.adhoc_source import codify_or_propose

# Max agent-invented (Ring-2) sources per run — bounds live network + LLM cost.
_RING2_MAX_PER_RUN = 5
from verticals.buyside_dd.comparator import compare_all
from verticals.buyside_dd.cross_claim_check import find_internal_divergences
from verticals.buyside_dd.policy import get_policy, RunPolicy
from verticals.buyside_dd.budget import BudgetState
from verticals.buyside_dd.materiality import infer_deal_context


HERE = Path(__file__).parent
INPUTS_DIR = HERE / "inputs"
OUTPUTS_DIR = HERE / "outputs"
FIXTURES_DIR = HERE / "fixtures"


def _run_tier1_discovery(claims: list[dict], deal_context: Optional[dict]) -> list[dict]:
    """Run Tier 1 open-web discovery on the issuer entity and convert any
    surfaced facts (addresses, related entities, etc.) into derived claims
    appended to the claim list for downstream verification.

    Returns derived-claim dicts in the same shape as LLM-extracted claims.
    """
    from verticals.buyside_dd.connectors.web_discovery import WebDiscoveryConnector
    from verticals.buyside_dd.connectors.base import ConnectorRequest

    # Find the issuer entity name. Prefer deal_context if it carries one;
    # otherwise grab the most-frequent entity-shaped subject across claims.
    issuer = None
    if deal_context:
        issuer = deal_context.get("issuer_name") or deal_context.get("entity_name")
    if not issuer:
        from collections import Counter
        subjects = Counter(c.get("subject") for c in claims if c.get("subject"))
        # The issuer is typically the subject of multiple claims
        for subj, n in subjects.most_common(5):
            if subj and len(subj) > 4 and any(c.isupper() for c in subj):
                issuer = subj
                break
    if not issuer:
        print("  [Tier 1] no issuer name resolvable from claims/context — skipping", file=sys.stderr)
        return []

    print(f"  [Tier 1] running web discovery on issuer: {issuer!r}", file=sys.stderr)

    # Three discovery queries, each tuned to a different fact category.
    # We accept partial success — if rate-limited on some queries, use what we get.
    query_recipes = [
        {"keywords": ["factory", "address"], "category": "physical_facility"},
        {"keywords": ["customers", "partners"], "category": "customer_pipeline"},
        {"keywords": ["headquarters", "office"], "category": "headquarters"},
    ]

    connector = WebDiscoveryConnector()
    derived = []
    for recipe in query_recipes:
        req = ConnectorRequest(entity_name=issuer, extra={"claim_keywords": recipe["keywords"]})
        try:
            r = connector.query(req)
        except Exception as e:
            print(f"    {recipe['category']}: connector exception ({e})", file=sys.stderr)
            continue
        if not r.success:
            print(f"    {recipe['category']}: {r.error_kind} ({(r.error_detail or '')[:60]})", file=sys.stderr)
            continue

        n_facts = 0
        for obs in r.observations:
            if obs.attribute == "discovery_meta":
                continue
            # Convert each observation to a derived claim
            if obs.attribute == "discovery_address":
                derived.append({
                    "claim_id": f"d_addr_{len(derived)+1}",
                    "subject": issuer,
                    "predicate": "located_at_address",
                    "object_value": obs.value,
                    "object_unit": None,
                    "scope": {"category": recipe["category"]},
                    "source_quote": (obs.extra or {}).get("snippet", "")[:200],
                    "source_doc": "tier1_web_discovery",
                    "discovery_url": obs.source_url,
                })
                n_facts += 1
            elif obs.attribute == "discovery_related_entity":
                derived.append({
                    "claim_id": f"d_ent_{len(derived)+1}",
                    "subject": issuer,
                    "predicate": "related_entity_named",
                    "object_value": obs.value,
                    "object_unit": None,
                    "scope": {"category": recipe["category"]},
                    "source_quote": (obs.extra or {}).get("snippet", "")[:200],
                    "source_doc": "tier1_web_discovery",
                    "discovery_url": obs.source_url,
                })
                n_facts += 1
            elif obs.attribute.startswith("discovery_url:"):
                cat = obs.attribute[len("discovery_url:"):]
                derived.append({
                    "claim_id": f"d_url_{len(derived)+1}",
                    "subject": issuer,
                    "predicate": f"surfaced_in_{cat}",
                    "object_value": obs.value,
                    "object_unit": None,
                    "scope": {"category": recipe["category"]},
                    "source_quote": (obs.extra or {}).get("title", "")[:200],
                    "source_doc": "tier1_web_discovery",
                    "discovery_url": obs.source_url,
                })
                n_facts += 1
        print(f"    {recipe['category']}: {n_facts} derived claims", file=sys.stderr)

    return derived


# ─── Stage 1: Document loading + claim extraction ────────────────────────────

SUPPORTED_DOC_TYPES = {".pdf", ".txt", ".md", ".html", ".csv", ".xlsx", ".docx", ".json"}


def _render_f_library_catalog() -> str:
    """Compact predicate catalog: predicate → description, sorted, deduped.
    The Mode B prompt feeds this to the LLM as the universe of verifiable
    predicates. Only predicates listed here are valid in Mode B output."""
    from verticals.buyside_dd.f_library import REAL_ESTATE_RULES
    by_pred: dict[str, list[str]] = {}
    for r in REAL_ESTATE_RULES:
        by_pred.setdefault(r.predicate, []).append(r.description)
    lines = []
    for pred in sorted(by_pred):
        descs = by_pred[pred]
        descs = list(dict.fromkeys(descs))[:2]
        lines.append(f"- `{pred}` — {'; '.join(descs)}")
    return "\n".join(lines)


def _render_source_atlas_summary() -> str:
    """List the IMPLEMENTED connectors (referent_type / attribute / source_id)
    so Mode B only generates verifications that can actually be dispatched.
    The full M-space taxonomy from source_atlas's docstring is broader than
    what's implemented — feeding the full taxonomy to the LLM produced Mode B
    claims using attribute keys we can't reach."""
    from verticals.buyside_dd.source_atlas import implemented_sources
    by_type: dict[str, list[tuple[str, str]]] = {}
    for s in implemented_sources():
        by_type.setdefault(s.referent_type.value, []).append((s.attribute, s.source_id))
    lines = []
    for rt in sorted(by_type):
        lines.append(f"\n{rt}:")
        attrs: dict[str, list[str]] = {}
        for a, sid in by_type[rt]:
            attrs.setdefault(a, []).append(sid)
        for a in sorted(attrs):
            sids = ", ".join(sorted(set(attrs[a])))
            lines.append(f"  - referent_attributes=['{a}']  via  {sids}")
    return "\n".join(lines)


def _parse_llm_json(response: str):
    """Strip markdown fences + any prose preamble before parsing JSON.

    Claude usually returns ```json [...] ``` or plain [...]. This helper
    handles both forms and falls back to extracting the first balanced
    {...} or [...] block if the response is wrapped in prose.
    """
    import re
    t = (response or "").strip()
    # Strip markdown fences like ```json\n...\n``` or ```\n...\n```
    if t.startswith("```"):
        t = re.sub(r"^```(?:json|JSON)?\s*", "", t)
        t = re.sub(r"\s*```\s*$", "", t)
    try:
        return json.loads(t)
    except json.JSONDecodeError:
        pass
    # Fallback: pull out the largest [...] or {...} block
    for opener, closer in (("[", "]"), ("{", "}")):
        i = t.find(opener)
        j = t.rfind(closer)
        if 0 <= i < j:
            try:
                return json.loads(t[i:j+1])
            except json.JSONDecodeError:
                continue
    # Last resort: re-raise the original error so the caller sees it
    return json.loads(t)


def discover_documents(input_path: Path) -> list[Path]:
    """Find all candidate documents to process."""
    if input_path.is_file():
        return [input_path]
    docs = []
    for p in sorted(input_path.rglob("*")):
        if p.is_file() and p.suffix.lower() in SUPPORTED_DOC_TYPES:
            docs.append(p)
    return docs


def load_document_text(doc_path: Path) -> str:
    """Convert a document to plain text for LLM processing.

    Real implementation should handle PDF (pypdf or pdfplumber), DOCX (python-docx),
    XLSX (openpyxl). For now, handles text-like formats and stubs the rest with
    a helpful error message.
    """
    suffix = doc_path.suffix.lower()
    if suffix in (".txt", ".md", ".csv", ".html", ".json"):
        try:
            return doc_path.read_text(encoding="utf-8", errors="ignore")
        except Exception as e:
            return f"[ERROR reading {doc_path}: {e}]"
    elif suffix == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(str(doc_path))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except ImportError:
            return f"[NEEDS pypdf: pip install pypdf — {doc_path}]"
        except Exception as e:
            return f"[ERROR reading PDF {doc_path}: {e}]"
    elif suffix == ".docx":
        try:
            import docx
            d = docx.Document(str(doc_path))
            return "\n\n".join(p.text for p in d.paragraphs)
        except ImportError:
            return f"[NEEDS python-docx: pip install python-docx — {doc_path}]"
        except Exception as e:
            return f"[ERROR reading DOCX {doc_path}: {e}]"
    elif suffix == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(doc_path), data_only=True)
            chunks = []
            for sheet in wb.sheetnames:
                ws = wb[sheet]
                chunks.append(f"=== Sheet: {sheet} ===")
                for row in ws.iter_rows(values_only=True):
                    chunks.append("\t".join(str(c) if c is not None else "" for c in row))
            return "\n".join(chunks)
        except ImportError:
            return f"[NEEDS openpyxl: pip install openpyxl — {doc_path}]"
        except Exception as e:
            return f"[ERROR reading XLSX {doc_path}: {e}]"
    return f"[UNSUPPORTED FORMAT {suffix}: {doc_path}]"


# ─── Stage 1 (continued): LLM claim extraction ───────────────────────────────

CLAIM_EXTRACTION_PROMPT = """You are extracting structured claims from an investment document for due diligence.

CANONICAL ISSUER (use this exact name as the `subject` for any claim about the company; do not use descriptive labels like "the company", "Factory 1", "[city] unit economics", etc.):
  {issuer_name}


Each "claim" is a discrete factual assertion the document makes that could in principle be verified against external evidence. Examples:
- "We acquired 1234 Main St on June 1 2023 for $500,000"
- "Current in-place rent for unit 2A is $1,800/mo"
- "Sponsor previously managed Fund III with 24% net IRR"
- "Property tax for 2024 will be $4,200"
- "Unit 5 was renovated in March 2024 with $35,000 in capex"

For each claim, output a JSON object with these fields:
- claim_id: a short unique identifier (e.g. "c001")
- subject: what the claim is about (a property address, sponsor entity, deal name, fund vintage, etc.)
- predicate: the assertion verb in snake_case (acquired_at_price, charges_rent, prior_fund_irr, projects_noi, has_lien, etc.)
- object_value: the asserted value
- object_unit: unit if applicable (USD, $/mo, %, ha, etc.)
- scope: dict with time_period, area, etc. as relevant
- source_quote: EXACT text from the document supporting this claim. Required. Used to detect hallucination.
- confidence: 0-1 your confidence that this is a real, intended claim from the doc

Only extract claims that are concrete and externally verifiable. Skip aspirational statements ("we believe..."), generic marketing ("best in class..."), or pure opinion. Skip internal commentary.

Return a JSON array. Aim for completeness — extract every distinct verifiable claim.

DOCUMENT TEXT:
---
{text}
---

Return ONLY the JSON array, no preamble or commentary."""


IMPLIED_VERIFICATIONS_PROMPT = """You are doing INVESTIGATION PLANNING for due diligence.

Stage A (extraction) already pulled the deck's own self-reported claims (Mode A). Your job is Stage B: derive ADDITIONAL VERIFICATIONS the deck IMPLIES but doesn't explicitly state. These are the claims an experienced DD analyst would add by reasoning about what should exist in external registries given what the deck says.

DEAL CONTEXT:
{deal_context}

CANONICAL ISSUER (use this exact name as the `subject` for any entity-level claim about the company itself; do not use descriptive labels):
  {issuer_name}

MODE A CLAIMS — what the deck says (summary; predicate, subject, value):
{mode_a_summary}

VERIFICATION TOOLS — f_library predicates the framework knows how to check:
{f_library_catalog}

IMPLEMENTED M-SOURCES (only these `referent_attributes` keys can actually be dispatched; do not use other attribute keys):
{source_atlas_summary}

INVESTIGATOR LOGIC — examples of Mode B derivations:
- Deck claims "we raised $7M Series A" → SEC Form D filing for the issuer SHOULD exist within 15 days of first sale (Reg D 506). Add a claim with predicate "raise_amount" and referent_attributes=["form_d_detail"] for the parent entity, even if the deck never mentions Form D.
- Deck claims "factory in Austin TX" → building permits SHOULD exist at the claimed address; OSHA establishment registry SHOULD have a record if production is ongoing. Add claims with predicate "operates_industrial_facility" pointing at city permit + OSHA sources.
- Deck claims "Antler led the round" → Antler SHOULD appear as a related person on the Form D OR have its own SAFE document. Add a claim with predicate "led_investment_round" subject=Antler.
- Deck claims "founder previously sold Tudu" → patent/trademark filings under founder's name SHOULD show that prior entity. Add a "previously_acquired" claim.
- Deck claims "10-year prior fund returned 24% IRR" → SEC IAPD / Form ADV SHOULD show that fund's existence. Add a "prior_fund_returned_irr" claim.

CRITICAL: each predicate you output MUST appear in the f_library catalog above. If a desirable verification doesn't have an f_library predicate, skip it — the framework can't run it.

For each implied verification, output a JSON object with the same schema as Mode A claims plus a `mode` field:
- claim_id: prefix "iv_" (e.g. "iv_form_d_parent", "iv_permit_austin")
- subject: the entity / property / person being verified
- predicate: MUST come from the f_library catalog above
- object_value: the asserted value (often the Mode A value, carried forward)
- object_unit: unit if applicable
- scope: dict with context (round name, address, expected_year, etc.)
- referent_type: from {{parcel, building, entity, person, event, amount, geographic_area, intangible}}.
  **CRITICAL: match the referent_type that the f-rule for your chosen predicate expects.** Predicates about a company, fund, or LLC (raise_amount, registered_as_entity, led_investment_round, filed_form_d_for_direct_safes, total_paid_to_AHC, valuation_cap, legal_status_is_clean, lien_status_is_clean) → `entity`. Predicates about a physical structure or specific address (located_at_address, operates_industrial_facility, building_condition_is_disclosed) → `building`. Predicates about a person (previously_acquired) → `person`. Predicates about market / area (projects_noi, projects_exit_cap_rate) → `geographic_area`. Predicates about a tax lot (acquired_at_price, property_market_value) → `parcel`. Do NOT use `event` for funding rounds — funding rounds are properties of the entity, not standalone events.
- referent_attributes: list of M-source attribute keys that should observe this claim. Must come from the IMPLEMENTED M-SOURCES list above.
- jurisdiction: state / city if known
- source_quote: 1-line investigator justification (NOT a quote from the deck — explain why this verification is implied)
- confidence: 1.0
- mode: "B"

Return ONLY a JSON array. Aim for 8-20 high-value implied verifications. Prioritize predicates with multiple f_library rules (highest leverage) and investigator-add-on verifications the deck doesn't explicitly call out.
"""


REFERENT_TYPING_PROMPT = """For each claim below, classify the REFERENT (what real-world thing the claim is about).

Referent types:
- parcel — a specific property identified by address or parcel number
- building — a specific physical structure (often colocated with parcel)
- entity — a company, LLC, fund, sponsor, government body
- person — a specific individual
- event — a transaction, transfer, payment, regulatory action
- amount — a generic financial quantity
- geographic_area — a region, market, neighborhood
- intangible — a license, certification, exemption, obligation

Also identify:
- referent_attributes: what about the referent the claim concerns (e.g., "sale_price", "current_rent", "prior_fund_returns")
- referent_id: any external ID hint (parcel #, BBL, address, LLC name)
- jurisdiction: state/city if identifiable

CLAIMS:
{claims_json}

Return a JSON array with one object per claim:
{{
  "claim_id": "...",
  "referent_type": "parcel" | "building" | etc.,
  "referent_attributes": ["..."],
  "referent_id": "..." | null,
  "jurisdiction": "..." | null,
  "verifiability": "high" | "medium" | "low" | "unverifiable"
}}"""


# ─── Pipeline driver ─────────────────────────────────────────────────────────

def run_pipeline(
    input_path: Path,
    output_dir: Path = OUTPUTS_DIR,
    llm_callable=None,
    policy: RunPolicy = None,
    deal_context: Optional[dict] = None,
):
    """
    Main pipeline. Pass llm_callable as a function (text, model="...") -> str
    to plug in an LLM (Claude API, etc.) for Stages 1-2.

    policy: RunPolicy (TRIAGE | STANDARD | DEEP). Defaults to STANDARD if None.
            Encodes budget caps, paid-source gating, and convergence rules.
    deal_context: optional dict with keys like 'deal_size_usd', 'asset_count' that
                  feed materiality derivation for deal-level claims.

    For environments without an LLM API, the pipeline returns
    intermediate state up to the point an LLM is needed, with a TODO marker.
    """
    if policy is None:
        policy = get_policy("standard")
    output_dir.mkdir(exist_ok=True, parents=True)
    run_id = datetime.utcnow().strftime("%Y%m%d_%H%M%S")
    run_dir = output_dir / run_id
    run_dir.mkdir(exist_ok=True)

    print(f"\n=== Buyside DD Pipeline — run {run_id} ===", file=sys.stderr)
    print(f"Input: {input_path}", file=sys.stderr)
    print(f"Output: {run_dir}", file=sys.stderr)

    # Stage 1a: Discover documents
    docs = discover_documents(input_path)
    print(f"\nStage 1a: Discovered {len(docs)} documents", file=sys.stderr)
    for d in docs:
        print(f"  - {d.relative_to(input_path) if input_path.is_dir() else d.name} ({d.stat().st_size:,} bytes)", file=sys.stderr)

    if not docs:
        print("\n[!] No documents found. Drop files in inputs/ and try again.", file=sys.stderr)
        return

    # Stage 1b: Load text
    print(f"\nStage 1b: Loading document text", file=sys.stderr)
    doc_texts = {}
    for d in docs:
        text = load_document_text(d)
        doc_texts[str(d)] = text
        warn = " [WARNING]" if text.startswith("[") else ""
        print(f"  - {d.name}: {len(text):,} chars{warn}", file=sys.stderr)

    # Save raw text for inspection
    with open(run_dir / "01_doc_texts.json", "w") as f:
        json.dump({k: v[:5000] for k, v in doc_texts.items()}, f, indent=2)

    # Stage 1c: Extract claims (LLM)
    print(f"\nStage 1c: Extracting claims via LLM", file=sys.stderr)
    if llm_callable is None:
        print("  [!] No LLM callable provided. Saving extraction prompts for manual run.", file=sys.stderr)
        prompts_dir = run_dir / "prompts"
        prompts_dir.mkdir(exist_ok=True)
        issuer_name_dry = (deal_context or {}).get("issuer_name") or "(unknown issuer)"
        for d, text in doc_texts.items():
            doc_name = Path(d).stem
            prompt = CLAIM_EXTRACTION_PROMPT.format(text=text[:50000], issuer_name=issuer_name_dry)
            (prompts_dir / f"{doc_name}_extraction.txt").write_text(prompt)
        print(f"  Saved {len(doc_texts)} extraction prompts to {prompts_dir}", file=sys.stderr)
        print(f"  → Manually feed each to an LLM and save results back as outputs/{run_id}/claims/{{doc_name}}.json", file=sys.stderr)
        return run_dir

    # If we have an LLM callable, run extraction
    issuer_name = (deal_context or {}).get("issuer_name") or "(unknown issuer; see scope.subject for entity name)"
    all_claims = []
    for d, text in doc_texts.items():
        doc_name = Path(d).stem
        prompt = CLAIM_EXTRACTION_PROMPT.format(text=text[:50000], issuer_name=issuer_name)
        try:
            response = llm_callable(prompt)
            claim_dicts = _parse_llm_json(response)
            for c in claim_dicts:
                c["source_doc"] = d
            all_claims.extend(claim_dicts)
            print(f"  - {d}: {len(claim_dicts)} claims", file=sys.stderr)
        except Exception as e:
            print(f"  - {d}: EXTRACTION FAILED ({e})", file=sys.stderr)

    with open(run_dir / "02_claims.json", "w") as f:
        json.dump(all_claims, f, indent=2, default=str)
    print(f"\nTotal claims extracted (Mode A): {len(all_claims)}", file=sys.stderr)

    if not all_claims:
        return run_dir

    # ── Stage 1c-B: Implied-verification derivation (Mode B) ──────────────
    # Good diligence has two modes. Mode A (above) extracts what the deck
    # SAYS — its own self-reported claims. Mode B derives what the deck
    # IMPLIES — verifications an investigator would run that the deck
    # doesn't explicitly mention (Form D existence, building permits at the
    # claimed address, Antler appearing as related person, etc.). See
    # memory/feedback_dd_two_modes.md for the framing.
    implied_claims: list[dict] = []
    if llm_callable is not None:
        print(f"\nStage 1c-B: Implied-verification derivation (Mode B)", file=sys.stderr)
        # Compact Mode A summary so the prompt stays under context
        mode_a_summary = "\n".join(
            f"- {c.get('predicate','?')} | subject={str(c.get('subject',''))[:60]} | "
            f"value={str(c.get('object_value',''))[:60]}"
            for c in all_claims[:80]
        )
        prompt = IMPLIED_VERIFICATIONS_PROMPT.format(
            deal_context=json.dumps(deal_context or {}, default=str),
            issuer_name=(deal_context or {}).get("issuer_name") or "(unknown — use the most-frequent entity-shaped subject from Mode A claims)",
            mode_a_summary=mode_a_summary,
            f_library_catalog=_render_f_library_catalog(),
            source_atlas_summary=_render_source_atlas_summary(),
        )
        # Save the prompt for audit (separate from the no-LLM dry-run
        # prompts dir, which is in a different code path)
        (run_dir / "02c_mode_b_prompt.txt").write_text(prompt)
        try:
            response = llm_callable(prompt)
            implied_claims = _parse_llm_json(response) or []
            for c in implied_claims:
                c.setdefault("source_doc", "mode_b_implied_verifications")
                c.setdefault("mode", "B")
            print(f"  Mode B derived {len(implied_claims)} implied verifications", file=sys.stderr)
            with open(run_dir / "02c_implied_verifications.json", "w") as f:
                json.dump(implied_claims, f, indent=2, default=str)
            all_claims.extend(implied_claims)
        except Exception as e:
            print(f"  Mode B FAILED: {e}", file=sys.stderr)

    # ── Stage 1.5: Entity Resolution ──────────────────────────────────────
    # The deck's stated subject is rarely the right query subject. A pitch
    # deck says "factory in Austin TX" but permits are keyed on the address,
    # the GC, and the landlord — none of which the deck names. Stage 1.5 runs
    # an LLM agent with web_search + registry tools to resolve deck subjects
    # into queryable real-world entities, with full evidence trails saved to
    # 02d_resolution_log.json. See entity_resolution.py for the architecture
    # and verticals/buyside_dd/RESOLUTION.md for the tool catalog.
    if llm_callable is not None:
        from verticals.buyside_dd.entity_resolution import (
            resolve_claims as _resolve_claims, get_resolver_client,
        )
        print(f"\nStage 1.5: Entity Resolution (LLM agent + tool use)", file=sys.stderr)
        resolver_client, resolver_model = get_resolver_client()
        if resolver_client is None:
            print(f"  No Anthropic client for resolver — skipping", file=sys.stderr)
        else:
            # The resolver needs a real-world anchor (issuer, city, state) to turn
            # relative deck subjects like "Factory 1" into queryable entities. The
            # late infer_deal_context() at the dispatch stage runs AFTER this, so
            # without an early inference the resolver sees an empty DEAL CONTEXT and
            # emits content-free queries (the AHC factory-resolution failure).
            # Operator-supplied context still wins on every key it sets.
            resolver_ctx = dict(deal_context or {})
            if not resolver_ctx.get("issuer_name"):
                early_ctx = infer_deal_context(all_claims)
                resolver_ctx = {**early_ctx, **resolver_ctx}
            new_children, resolution_log = _resolve_claims(
                claims=all_claims,
                deal_context=resolver_ctx,
                client=resolver_client,
                model=resolver_model,
                max_turns_per_claim=6,
            )
            with open(run_dir / "02d_resolution_log.json", "w") as f:
                json.dump(resolution_log, f, indent=2, default=str)
            if new_children:
                with open(run_dir / "02d_resolved_children.json", "w") as f:
                    json.dump(new_children, f, indent=2, default=str)
                all_claims.extend(new_children)
            print(f"  Stage 1.5: {len(resolution_log)} claims processed, "
                  f"{len(new_children)} resolved child claims appended",
                  file=sys.stderr)

    # ── Stage 1d: Tier 1 open-web discovery ───────────────────────────────
    # The deck is the source of R, not an exhaustive list of what's true.
    # Before running Tier 2/3 verifiers (which need specific facts to query),
    # ask the open web what it knows about the issuer entity. Surfaces
    # addresses, partners, customers, GCs, etc. that aren't in the deck and
    # gives downstream verifiers concrete things to test.
    print(f"\nStage 1d: Tier 1 open-web discovery", file=sys.stderr)
    discovery_claims = _run_tier1_discovery(all_claims, deal_context)
    if discovery_claims:
        all_claims.extend(discovery_claims)
        with open(run_dir / "02b_discovery_claims.json", "w") as f:
            json.dump(discovery_claims, f, indent=2, default=str)
        print(f"  Surfaced {len(discovery_claims)} derived claims from web discovery", file=sys.stderr)
    else:
        print(f"  No derived claims surfaced (web discovery returned nothing usable)", file=sys.stderr)

    # Stage 2: Type referents (LLM)
    print(f"\nStage 2: Typing referents", file=sys.stderr)
    typing_prompt = REFERENT_TYPING_PROMPT.format(claims_json=json.dumps(all_claims, indent=2))
    try:
        typing_response = llm_callable(typing_prompt)
        typings = _parse_llm_json(typing_response)
        # Merge typings back into claims, but Mode B claims already carry a
        # referent_type chosen against the f_library catalog — don't let the
        # typing stage clobber it (the typing prompt doesn't see f_library
        # and will classify e.g. raise_amount claims as "event" when
        # raise_amount's f-rule expects "entity").
        typing_by_id = {t["claim_id"]: t for t in typings}
        typed_claims = []
        for c in all_claims:
            t = typing_by_id.get(c["claim_id"])
            if t:
                if c.get("mode") == "B" and c.get("referent_type"):
                    # Preserve Mode B's referent_type + referent_attributes;
                    # let typing fill in jurisdiction / verifiability if it
                    # has stronger signal there.
                    merged = {**t, **c}
                else:
                    merged = {**c, **t}
                typed_claims.append(merged)
            else:
                typed_claims.append({**c, "referent_type": c.get("referent_type") or "unknown"})
        with open(run_dir / "03_typed_claims.json", "w") as f:
            json.dump(typed_claims, f, indent=2, default=str)
        print(f"  Typed {len(typed_claims)} claims", file=sys.stderr)
    except Exception as e:
        print(f"  TYPING FAILED ({e})", file=sys.stderr)
        return run_dir

    # Stage 3: f-rule lookup
    print(f"\nStage 3: f-rule lookup", file=sys.stderr)
    claims_with_f = []
    n_with_f = 0
    for c in typed_claims:
        ref_type_str = c.get("referent_type", "unknown")
        try:
            ref_type = ReferentType(ref_type_str)
        except ValueError:
            ref_type = ReferentType.UNKNOWN
        rules = find_rules(c["predicate"], ref_type, c.get("jurisdiction"))
        c["f_rules"] = [r.model_dump() for r in rules]
        c["f_rule_count"] = len(rules)
        if rules:
            n_with_f += 1
        claims_with_f.append(c)
    print(f"  {n_with_f}/{len(typed_claims)} claims matched at least one f-rule", file=sys.stderr)
    with open(run_dir / "04_claims_with_f.json", "w") as f:
        json.dump(claims_with_f, f, indent=2, default=str)

    # Stage 4: M source lookup
    print(f"\nStage 4: M source lookup", file=sys.stderr)
    claims_with_sources = []
    n_with_sources = 0
    n_recall_floor_added = 0
    n_ring2_invented = 0
    for c in claims_with_f:
        ref_type = ReferentType(c.get("referent_type", "unknown"))
        sources = []
        for attr in c.get("referent_attributes", []):
            sources.extend(find_sources(ref_type, attr, c.get("jurisdiction")))
        # Dedupe by source_id
        seen_ids = set()
        unique_sources = []
        for s in sources:
            if s.source_id not in seen_ids:
                seen_ids.add(s.source_id)
                unique_sources.append(s)
        # Ring-1 recall floor: a mandatory source the agent OMITTED (e.g. OSHA on
        # a factory). Surface loudly AND enforce by adding the expected source so
        # the silent skip can't happen.
        selected_ids = [s.source_id for s in unique_sources]
        gaps = recall_floor_gaps(c, selected_ids)
        for g in gaps:
            for sid in g["missing_source_ids"]:
                src = source_by_id(sid)
                if src and src.source_id not in seen_ids:
                    seen_ids.add(src.source_id)
                    unique_sources.append(src)
                    n_recall_floor_added += 1
            print(f"  [recall-floor] {c.get('claim_id')} ({c.get('subject','')[:30]}): "
                  f"feature={g['feature']} → added {g['missing_source_ids']} "
                  f"(agent omitted; {g['why']})", file=sys.stderr)
        # Ring-2 open-set escape hatch: the claim names a verifiable attribute but
        # NO registered (Ring-0/1) source covers it. Rather than silently dropping
        # it, let the agent invent a public read-only source, run it sandboxed, and
        # propose it back to the atlas (gated). Findings from it are PROVISIONAL.
        if (not unique_sources and c.get("referent_attributes")
                and llm_callable is not None and n_ring2_invented < _RING2_MAX_PER_RUN):
            for attr in c.get("referent_attributes", []):
                if n_ring2_invented >= _RING2_MAX_PER_RUN:
                    break
                try:
                    summary = codify_or_propose(
                        referent_type=c.get("referent_type", "unknown"),
                        attribute=attr,
                        subject=c.get("subject", ""),
                        scope=c.get("scope") or {},
                        llm=llm_callable,
                        run_dir=run_dir,
                    )
                except Exception as e:
                    print(f"  [ring2] invention failed for {c.get('claim_id')}/{attr}: {e}", file=sys.stderr)
                    continue
                if not summary:
                    continue
                n_ring2_invented += 1
                c.setdefault("ring2_provisional", []).append(summary)
                print(f"  [ring2] {c.get('claim_id')} ({c.get('subject','')[:24]}) invented "
                      f"'{summary['source_id']}' ran_ok={summary['ran_ok']} "
                      f"gate={summary['promotion_gate']} → {summary['artifact']}", file=sys.stderr)
        c["m_sources"] = [s.model_dump() for s in unique_sources]
        c["m_source_count"] = len(unique_sources)
        if unique_sources:
            n_with_sources += 1
        claims_with_sources.append(c)
    print(f"  {n_with_sources}/{len(claims_with_f)} claims have at least one M source"
          + (f" (+{n_recall_floor_added} added by recall floor)" if n_recall_floor_added else "")
          + (f"; {n_ring2_invented} Ring-2 provisional sources invented" if n_ring2_invented else ""),
          file=sys.stderr)
    with open(run_dir / "05_claims_with_sources.json", "w") as f:
        json.dump(claims_with_sources, f, indent=2, default=str)

    # Auto-infer deal_context if not provided by operator
    if not deal_context:
        deal_context = infer_deal_context(claims_with_sources)
        print(f"\n[deal_context inferred] asset_count={deal_context.get('asset_count')} "
              f"deal_size_usd={deal_context.get('deal_size_usd')} mode={deal_context.get('mode')}", file=sys.stderr)
    else:
        # Operator passed some context; backfill missing fields via inference
        inferred = infer_deal_context(claims_with_sources)
        for k, v in inferred.items():
            deal_context.setdefault(k, v)
        print(f"\n[deal_context] {deal_context}", file=sys.stderr)

    # Stage 5: Policy-driven acquisition
    print(f"\nStage 5: Policy-driven dispatch ({policy.name}, max_cost=${policy.max_cost_usd}, paid={policy.paid_sources_enabled})", file=sys.stderr)
    cws_objects = []
    for c in claims_with_sources:
        try:
            ref_type = ReferentType(c.get("referent_type", "unknown"))
        except ValueError:
            ref_type = ReferentType.UNKNOWN
        claim = Claim(
            claim_id=c["claim_id"],
            source_doc=c.get("source_doc", ""),
            source_quote=c.get("source_quote", ""),
            subject=c.get("subject", ""),
            predicate=c.get("predicate", ""),
            object_value=c.get("object_value"),
            object_unit=c.get("object_unit"),
            scope=c.get("scope") or {},
            confidence=c.get("confidence", 1.0),
        )
        tc = TypedClaim(
            claim=claim,
            referent_type=ref_type,
            referent_attributes=c.get("referent_attributes", []),
            referent_id=c.get("referent_id"),
            jurisdiction=c.get("jurisdiction"),
        )
        rules = [FRule(**r) for r in c.get("f_rules", [])]
        sources = [MSource(**s) for s in c.get("m_sources", [])]
        cws = ClaimWithSources(typed_claim=tc, f_rules=rules, m_sources=sources)
        cws_objects.append((cws, c))

    cov = coverage_report([x[0] for x in cws_objects])
    print(f"  Coverage: {cov['claims_with_implemented_connector']}/{cov['total_claims']} claims have an implemented connector ({cov['implementation_coverage_pct']}%)", file=sys.stderr)

    # Single budget across the whole run
    budget = BudgetState(max_cost_usd=policy.max_cost_usd, max_wall_time_min=policy.max_wall_time_min)

    all_findings = []
    n_queried = 0
    n_deferred = 0
    n_immaterial = 0
    n_budget_hit = 0
    for cws, raw_dict in cws_objects:
        if not cws.m_sources or not cws.f_rules:
            continue
        outcome = dispatch_with_policy(cws, policy=policy, budget=budget, deal_context=deal_context)
        n_queried += 1
        if outcome.stop_reason.value == "deferred":
            n_deferred += 1
        elif outcome.stop_reason.value == "immaterial":
            n_immaterial += 1
        elif outcome.stop_reason.value == "budget_hit":
            n_budget_hit += 1
        successes = sum(1 for r in outcome.results if r.success)
        print(f"  - {cws.typed_claim.claim.claim_id} ({cws.typed_claim.claim.subject[:30]}): "
              f"{successes}/{len(outcome.results)} sources OK; "
              f"stop={outcome.stop_reason.value}; mat={outcome.materiality.status}", file=sys.stderr)
        findings = compare_all(cws.typed_claim, cws.f_rules, outcome, cws.m_sources)
        all_findings.extend(findings)

    # Stage 5b: Cross-claim internal-consistency check (Mode A). Needs no
    # external M — divergences come from the materials contradicting themselves
    # (e.g. factory capacity 1,000 homes/yr in the deck vs 1,750 in the model).
    print(f"\nStage 5b: Cross-claim consistency check", file=sys.stderr)
    internal_findings = find_internal_divergences([cws.typed_claim for cws, _ in cws_objects])
    for fd in internal_findings:
        print(f"  [{fd.severity.value.upper():11s}] {fd.f_rule_id} on "
              f"'{fd.claim.subject[:30]}' / '{fd.claim.predicate}' — "
              f"divergence={(fd.divergence_pct or 0)*100:.1f}%", file=sys.stderr)
    all_findings.extend(internal_findings)

    with open(run_dir / "06_findings.json", "w") as f:
        json.dump([fd.model_dump() for fd in all_findings], f, indent=2, default=str)
    with open(run_dir / "06_budget.json", "w") as f:
        json.dump(budget.summary(), f, indent=2, default=str)

    # Stage 6: Severity rollup
    print(f"\nStage 6: Severity rollup", file=sys.stderr)
    by_severity = {}
    for fd in all_findings:
        by_severity.setdefault(fd.severity.value, []).append(fd)
    for sev in ("critical", "severe", "moderate", "minor", "pass", "unverifiable"):
        ct = len(by_severity.get(sev, []))
        if ct:
            print(f"  {sev.upper()}: {ct}", file=sys.stderr)

    # Build the final report
    report = {
        "run_id": run_id,
        "timestamp": datetime.utcnow().isoformat(),
        "input": str(input_path),
        "policy": {"name": policy.name, "version": policy.version},
        "deal_context": deal_context or {},
        "n_documents": len(docs),
        "n_claims_extracted": len(all_claims),
        "n_claims_typed": len(typed_claims),
        "n_claims_with_f_rule": n_with_f,
        "n_claims_with_m_source": n_with_sources,
        "n_claims_dispatched": n_queried,
        "n_findings": len(all_findings),
        "n_deferred_to_human": n_deferred,
        "n_skipped_immaterial": n_immaterial,
        "n_stopped_budget": n_budget_hit,
        "coverage": cov,
        "budget": budget.summary(),
        "severity_counts": {sev: len(fs) for sev, fs in by_severity.items()},
        "stop_reason_counts": _count_stop_reasons(all_findings),
        "total_exposure_usd": sum(
            fd.exposure_estimate_usd or 0 for fd in all_findings
            if fd.severity in (Severity.MODERATE, Severity.SEVERE, Severity.CRITICAL)
        ),
        "stage_status": {
            "1_extract": "ok" if all_claims else "no_claims",
            "2_type": "ok" if typed_claims else "failed",
            "3_f_lookup": "ok",
            "4_source_lookup": "ok",
            "5_acquisition": "ok",
            "6_compare": "ok",
        },
    }
    with open(run_dir / "07_report.json", "w") as f:
        json.dump(report, f, indent=2, default=str)

    print(f"\n=== Pipeline complete ===", file=sys.stderr)
    print(f"Findings: {len(all_findings)} | Exposure (mod+sev+crit): ${report['total_exposure_usd']:,.0f}", file=sys.stderr)
    print(f"Deferred to human: {n_deferred} | Immaterial: {n_immaterial} | Budget hit: {n_budget_hit}", file=sys.stderr)
    print(f"Budget: ${budget.summary()['spent_usd']}/{budget.summary()['max_cost_usd']} | calls: {budget.summary()['calls_made']}", file=sys.stderr)
    print(f"Output: {run_dir}", file=sys.stderr)
    return run_dir


def _count_stop_reasons(findings):
    counts = {}
    for fd in findings:
        sr = fd.stop_reason.value
        counts[sr] = counts.get(sr, 0) + 1
    return counts


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(f"Usage: python3 pipeline.py <input_path> [--policy=triage|standard|deep] [--deal-size=USD]", file=sys.stderr)
        print(f"  input_path: file or directory containing DD materials", file=sys.stderr)
        sys.exit(1)
    input_path = Path(sys.argv[1])
    if not input_path.exists():
        print(f"Input not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    policy_name = "standard"
    deal_context = {}
    for arg in sys.argv[2:]:
        if arg.startswith("--policy="):
            policy_name = arg.split("=", 1)[1]
        elif arg.startswith("--deal-size="):
            try:
                deal_context["deal_size_usd"] = float(arg.split("=", 1)[1])
            except ValueError:
                pass
        elif arg.startswith("--issuer="):
            deal_context["issuer_name"] = arg.split("=", 1)[1]
    chosen_policy = get_policy(policy_name)
    print(f"[policy: {chosen_policy.name}; deal_context: {deal_context}]", file=sys.stderr)

    # Try to set up Anthropic LLM if API key present
    llm = None
    import os
    if os.getenv("ANTHROPIC_API_KEY"):
        try:
            import anthropic
            client = anthropic.Anthropic()
            def llm(prompt: str, model: str = "claude-sonnet-4-5"):
                # Stream so we can request a high max_tokens without the
                # SDK's >10-min guard rejecting the call. Aggregate text
                # deltas back into a single string for the caller.
                parts: list[str] = []
                with client.messages.stream(
                    model=model,
                    max_tokens=32000,
                    messages=[{"role": "user", "content": prompt}],
                ) as stream:
                    for chunk in stream.text_stream:
                        parts.append(chunk)
                return "".join(parts)
            print("[llm: Anthropic API configured]", file=sys.stderr)
        except ImportError:
            print("[llm: anthropic package not installed; pip install anthropic]", file=sys.stderr)
        except Exception as e:
            print(f"[llm: Anthropic setup failed: {e}]", file=sys.stderr)

    run_pipeline(input_path, llm_callable=llm, policy=chosen_policy, deal_context=deal_context)
