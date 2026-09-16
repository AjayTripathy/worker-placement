"""lca_pyramid — DOL H-1B/LCA multi-quarter fetcher + the PYRAMID nowcast for IT-services names.

Born 2026-07-09 (CTSH DD): headcount IS a consultancy's P&L, and DOL publishes every LCA
(job title, wage, PW WAGE LEVEL I-IV, worksite) quarterly — a physical-truth labor lead the
sell-side doesn't systematically track. The single-quarter snapshot supported "pyramid
flattening" but the TREND was a coverage gap (no fetcher, one quarter on disk). This closes it.

The signal (per employer, quarter over quarter):
  - VOLUME (# LCAs) + median annualized WAGE + WAGE-LEVEL mix (I entry .. IV expert).
  - PYRAMID FLATTENING  = volume DOWN + wage/level UP  -> pruning the automatable junior base,
    keeping high-value work = the AI-adaptation SURVIVAL path (bullish-ish for a services name).
  - CONTRACTION         = volume DOWN + wage/level DOWN/flat -> broad demand loss = the TRAP.
  - EXPANSION           = volume UP.
  Directly feeds the CTSH-type kill trigger: "LCA volume AND wage-level both falling".

Pipeline (gov source is Akamai-gated, so fetch is best-effort with a manual fallback):
  python3 -m verticals.public_co.lca_pyramid fetch FY2026_Q2   # try download -> slim, else instruct
  python3 -m verticals.public_co.lca_pyramid slim <raw.xlsx|.csv> FY2026_Q2   # slim a placed file
  python3 -m verticals.public_co.lca_pyramid trend             # the nowcast across all quarters
READ-ONLY on markets. Raw files are gitignored; only the small target-filtered slims persist.
"""
from __future__ import annotations

import csv
import statistics
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "verticals" / "public_co" / "data" / "_dol_data"          # huge raw DOL files (gitignored)
DATA = ROOT / "verticals" / "public_co" / "data" / "lca_pyramid"        # small target slims + state (tracked)
STATE = DATA / "lca_pyramid_state.json"
DOL_BASE = "https://www.dol.gov/sites/dolgov/files/ETA/oflc/pdfs"

# IT-services / consulting employers we track -> ticker. Name match is UPPER-cased substring.
TARGET_EMPLOYERS = {
    "COGNIZANT": "CTSH", "INFOSYS": "INFY", "WIPRO": "WIT", "TATA CONSULTANCY": "TCS",
    "ACCENTURE": "ACN", "EPAM": "EPAM", "GENPACT": "G", "CONCENTRIX": "CNXC",
    "HCL": "HCLTECH", "CAPGEMINI": "CAP", "TECH MAHINDRA": "TECHM", "DXC": "DXC",
    "IBM": "IBM", "MPHASIS": "MPHASIS", "LTIMINDTREE": "LTIM",
}
WAGE_MULT = {"YEAR": 1, "HOUR": 2080, "MONTH": 12, "WEEK": 52, "BI-WEEKLY": 26, "BIWEEKLY": 26}


def _ticker(emp: str):
    u = (emp or "").upper()
    for frag, tk in TARGET_EMPLOYERS.items():
        if frag in u:
            return tk
    return None


def _annual_wage(raw, unit):
    try:
        w = float(str(raw).replace(",", "").replace("$", "").strip())
    except Exception:
        return None
    return w * WAGE_MULT.get((unit or "YEAR").upper().strip(), 1)


def _colmap(header):
    """Fuzzy-map the DOL raw/slim header (names drift across FY vintages) to our fields."""
    h = [c.strip().upper() for c in header]

    def find(*subs, avoid=()):
        for i, c in enumerate(h):
            if all(s in c for s in subs) and not any(a in c for a in avoid):
                return i
        return None
    return {
        "employer": find("EMPLOYER", "NAME", avoid=("AGENT", "ATTORNEY", "SECONDARY")) or find("EMPLOYER_NAME") or find("EMPLOYER"),
        "status": find("CASE", "STATUS") or find("STATUS"),
        "visa": find("VISA", "CLASS"),
        "title": find("JOB", "TITLE"),
        "soc": find("SOC", "TITLE"),
        "wage": find("WAGE", "FROM") or find("WAGE_RATE") or find("WAGE_FROM"),
        "unit": find("WAGE", "UNIT") or find("UNIT", "PAY"),
        "level": find("PW", "LEVEL") or find("WAGE_LEVEL") or find("PW_WAGE_LEVEL"),
        "state": find("WORKSITE", "STATE"),
        "pos": find("TOTAL", "WORKER", "POSITIONS") or find("TOTAL_WORKER_POSITIONS"),
    }


def _iter_rows(path: Path):
    """Yield rows (list) from an xlsx (streamed) or a csv. Header first."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(filename=str(path), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(values_only=True):
            yield list(row)
        wb.close()
    else:
        with open(path, newline="") as f:
            for r in csv.reader(f):
                yield r


def slim(raw_path, quarter):
    """Filter a raw DOL LCA file (or an existing slim CSV) to TARGET_EMPLOYERS -> small per-quarter CSV."""
    raw = Path(raw_path)
    assert raw.exists(), f"file not found: {raw}"
    out = DATA / f"lca_{quarter}_targets.csv"
    DATA.mkdir(parents=True, exist_ok=True)
    rows = _iter_rows(raw)
    header = next(rows)
    cm = _colmap(header)
    if cm["employer"] is None:
        raise SystemExit(f"could not locate EMPLOYER column in header: {header[:12]}...")
    n_in = n_kept = 0
    with open(out, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["ticker", "quarter", "employer", "status", "visa", "title", "soc",
                    "wage_annual", "wage_level", "state", "positions"])
        for r in rows:
            n_in += 1
            emp = r[cm["employer"]] if cm["employer"] < len(r) else ""
            tk = _ticker(emp)
            if not tk:
                continue
            def g(k):
                i = cm[k]
                return r[i] if (i is not None and i < len(r)) else ""
            wa = _annual_wage(g("wage"), g("unit"))
            w.writerow([tk, quarter, emp, g("status"), g("visa"), g("title"), g("soc"),
                        round(wa) if wa else "", g("level"), g("state"), g("pos")])
            n_kept += 1
    print(f"[lca] slimmed {quarter}: {n_in:,} rows -> {n_kept:,} target-employer rows -> {out.name}")
    return out


def _load_all():
    """All target slims -> {ticker: {quarter: [rows]}}."""
    by = defaultdict(lambda: defaultdict(list))
    for p in sorted(DATA.glob("lca_*_targets.csv")):
        for row in csv.DictReader(open(p)):
            by[row["ticker"]][row["quarter"]].append(row)
    return by


def _qtr_stats(rows):
    certified = [r for r in rows if "CERTIF" in (r.get("status") or "").upper()]
    wages = [float(r["wage_annual"]) for r in rows if r.get("wage_annual")]
    levels = [str(r.get("wage_level") or "").strip().upper() for r in rows]
    lv = {k: levels.count(k) for k in ("LEVEL I", "LEVEL II", "LEVEL III", "LEVEL IV")}
    # tolerate bare 'I'/'1' encodings
    for r in rows:
        s = str(r.get("wage_level") or "").strip().upper()
        for k, alts in {"LEVEL I": ("I", "1"), "LEVEL II": ("II", "2"), "LEVEL III": ("III", "3"), "LEVEL IV": ("IV", "4")}.items():
            if s in alts:
                lv[k] += 1
    has_level = sum(lv.values()) > 0
    med = statistics.median(wages) if wages else None
    # junior% on a WAGE-TIER basis ALWAYS (comparable across quarters even when one lacks the level
    # column) — the level mix below is display-only, NOT used for the cross-quarter delta.
    junior_pct = (sum(1 for w in wages if w < 95000) / len(wages)) if wages else None
    return {"n": len(rows), "certified": len(certified), "median_wage": med,
            "levels": lv, "has_level": has_level, "junior_pct": junior_pct}


def trend():
    """The nowcast: per ticker, quarter-over-quarter volume + wage/level -> pyramid classification."""
    import json
    by = _load_all()
    if not by:
        print("[lca] no target slims yet — run `fetch <FYQ>` or `slim <file> <FYQ>` first.")
        return
    out = {}
    print("=" * 74)
    print("DOL H-1B/LCA PYRAMID NOWCAST  —  headcount is the consultancy's P&L")
    print("junior% = wage-tier <$95k (comparable across quarters); L I/II/III/IV mix shown where present.")
    print("VOLUME + median WAGE are the robust deltas; small US-visa books (EPAM/G/CNXC, n<70) are NOISY.")
    print("=" * 74)
    for tk in sorted(by):
        quarters = sorted(by[tk])
        stats = {q: _qtr_stats(by[tk][q]) for q in quarters}
        latest = quarters[-1]
        s = stats[latest]
        line = f"\n{tk}  [{len(quarters)}q: {quarters[0]}..{latest}]"
        print(line)
        for q in quarters:
            st = stats[q]
            jp = f"{st['junior_pct']*100:.0f}%" if st["junior_pct"] is not None else "—"
            mw = f"${st['median_wage']:,.0f}" if st["median_wage"] else "—"
            lvl = ("  L I/II/III/IV " + "/".join(str(st["levels"][k]) for k in ("LEVEL I", "LEVEL II", "LEVEL III", "LEVEL IV"))) if st["has_level"] else "  (no level col)"
            print(f"   {q}: n={st['n']:>5}  cert={st['certified']:>5}  med={mw:>10}  junior={jp:>4}{lvl}")
        verdict = "SNAPSHOT (need >=2 quarters for a trajectory)"
        if len(quarters) >= 2:
            prev, cur = stats[quarters[-2]], s
            dv = (cur["n"] - prev["n"]) / max(1, prev["n"])
            dw = ((cur["median_wage"] or 0) - (prev["median_wage"] or 0)) / max(1, prev["median_wage"] or 1) if prev["median_wage"] else 0
            dj = (cur["junior_pct"] or 0) - (prev["junior_pct"] or 0) if prev["junior_pct"] is not None else 0
            if dv > 0.10:
                verdict = f"EXPANSION (volume {dv:+.0%})"
            elif dv < -0.05 and (dw > 0.02 or dj < -0.02):
                verdict = f"PYRAMID FLATTENING (vol {dv:+.0%}, wage {dw:+.0%}, junior {dj*100:+.0f}pp) — pruning juniors, keeping value"
            elif dv < -0.05 and dw <= 0.02:
                verdict = f"CONTRACTION (vol {dv:+.0%}, wage {dw:+.0%}) — broad demand loss (the TRAP signal)"
            else:
                verdict = f"STABLE (vol {dv:+.0%}, wage {dw:+.0%})"
        print(f"   -> {verdict}")
        out[tk] = {"quarters": quarters, "latest": latest, "verdict": verdict,
                   "latest_stats": {k: v for k, v in s.items() if k != "levels"}}
    STATE.write_text(json.dumps(out, indent=1, default=str))
    print(f"\n[lca] -> {STATE.name}")
    return out


PERF_PAGE = "https://www.dol.gov/agencies/eta/foreign-labor/performance"
_UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"


def _playwright_fetch(url, dest):
    """Land on the DOL performance page (clears Akamai, sets cookies), then stream the file down —
    the proven gated-gov-portal pattern (real Chrome, land-first, in-context fetch)."""
    try:
        from playwright.sync_api import sync_playwright
    except Exception as e:
        return False, f"playwright import: {e}"
    ok, err = False, ""
    with sync_playwright() as p:
        b = p.chromium.launch(channel="chrome", headless=True)
        ctx = b.new_context(user_agent=_UA, accept_downloads=True)
        pg = ctx.new_page()
        try:
            pg.goto(PERF_PAGE, wait_until="domcontentloaded", timeout=60000)
            pg.wait_for_timeout(3500)                       # let the anti-bot challenge clear
            try:                                            # primary: streamed navigation download (no 300MB in RAM)
                with pg.expect_download(timeout=900000) as dl:
                    try:
                        pg.goto(url, timeout=900000)
                    except Exception:
                        pass                                # ERR_ABORTED on a download is expected
                dl.value.save_as(str(dest))
                ok = dest.exists() and dest.stat().st_size > 1_000_000
            except Exception:
                ok = False
            if not ok:                                      # fallback: context request reuses the cleared cookies
                resp = ctx.request.get(url, timeout=900000)
                if resp.ok:
                    dest.write_bytes(resp.body())
                    ok = dest.exists() and dest.stat().st_size > 1_000_000
                else:
                    err = f"context.request HTTP {resp.status}"
        except Exception as e:
            err = f"{type(e).__name__}: {e}"
        finally:
            b.close()
    return ok, err


def fetch(quarter):
    """Auto-download a DOL LCA quarter (Playwright real-Chrome past the Akamai gate), then slim.
    Falls back to explicit manual instructions if the browser fetch fails."""
    RAW.mkdir(parents=True, exist_ok=True)
    url = f"{DOL_BASE}/LCA_Disclosure_Data_{quarter}.xlsx"
    raw = RAW / f"LCA_Disclosure_Data_{quarter}.xlsx"
    print(f"[lca] Playwright fetch (land-first past Akamai): {url}")
    ok, err = _playwright_fetch(url, raw)
    if ok:
        print(f"[lca] downloaded {raw.stat().st_size/1e6:.0f}MB — slimming…")
        slim(raw, quarter)
        raw.unlink()                                        # keep only the small target slim
        print(f"[lca] {quarter} ingested. Run `trend` for the updated nowcast.")
        return
    if raw.exists():
        raw.unlink()
    print(f"[lca] auto-fetch failed ({err or 'blocked'}). MANUAL PATH:")
    print(f"      1. In a browser, download: {url}")
    print(f"      2. Save it to: {RAW}/")
    print(f"      3. Run: python3 -m verticals.public_co.lca_pyramid slim {RAW}/LCA_Disclosure_Data_{quarter}.xlsx {quarter}")


import datetime


def _fyq(d):
    """Calendar date -> DOL fiscal (FY, quarter). FY starts Oct; Q1=Oct-Dec .. Q4=Jul-Sep."""
    m, y = d.month, d.year
    if m >= 10:
        return (y + 1, 1)
    if m <= 3:
        return (y, 2)
    if m <= 6:
        return (y, 3)
    return (y, 4)


def _dec(fy, q):
    q -= 1
    if q == 0:
        q, fy = 4, fy - 1
    return fy, q


def _recent(n=6, asof=None):
    fy, q = _fyq(asof or datetime.date.today())
    fy, q = _dec(fy, q)                      # skip the CURRENT in-progress quarter (never published)
    out = []
    for _ in range(n):
        out.append(f"FY{fy}_Q{q}")
        fy, q = _dec(fy, q)
    return out


def _have():
    return {p.name.replace("lca_", "").replace("_targets.csv", "") for p in DATA.glob("lca_*_targets.csv")}


def auto():
    """Full automation: fetch any recent completed quarter we don't yet have, then recompute the nowcast.
    Idempotent + self-healing — an unpublished/blocked quarter fails gracefully and retries next run."""
    have = _have()
    missing = [q for q in _recent(6) if q not in have]
    if missing:
        print(f"[lca auto] have {sorted(have)}; fetching missing: {missing}")
        for q in missing:
            fetch(q)
    else:
        print(f"[lca auto] all recent quarters present ({sorted(have)}).")
    trend()


def main():
    a = sys.argv[1:]
    if not a:
        print(__doc__)
        return
    if a[0] == "fetch" and len(a) >= 2:
        fetch(a[1])
    elif a[0] == "slim" and len(a) >= 3:
        slim(a[1], a[2])
    elif a[0] == "trend":
        trend()
    elif a[0] == "auto":
        auto()
    else:
        print("usage: auto | fetch <FYQ> | slim <file> <FYQ> | trend")


if __name__ == "__main__":
    main()
