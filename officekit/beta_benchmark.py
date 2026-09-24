"""Dated public benchmark snapshot. Only an explicit refresh performs I/O."""
from datetime import datetime
from io import BytesIO
from urllib.request import Request, urlopen

from officekit.beta_programs import normalize_snapshot

URL = 'https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/holdings-daily-us-en-sptm.xlsx'


def fetch():
    from openpyxl import load_workbook
    request = Request(URL, headers={'User-Agent': 'Mozilla/5.0 WorkerPlacement benchmark review'})
    with urlopen(request, timeout=30) as response:
        data = response.read(8 * 1024 * 1024 + 1)
    if len(data) > 8 * 1024 * 1024:
        raise ValueError('The benchmark download exceeds the size limit.')
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    raw_date = next((r[1] for r in rows[:12] if r and r[0] == 'Holdings:'), None)
    if isinstance(raw_date, datetime):
        as_of = raw_date.date().isoformat()
    else:
        value = str(raw_date).strip().removeprefix('As of ').strip()
        as_of = None
        for fmt in ('%d-%b-%Y', '%Y-%m-%d', '%m/%d/%Y', '%d %b %Y'):
            try:
                as_of = datetime.strptime(value, fmt).date().isoformat()
                break
            except ValueError:
                pass
        if not as_of:
            raise ValueError('The benchmark publication date could not be read.')
    index = next(i for i, r in enumerate(rows) if r and r[0] == 'Name' and 'Ticker' in r and 'Weight' in r)
    headers = {name: i for i, name in enumerate(rows[index])}
    constituents = []
    for r in rows[index + 1:]:
        s, w = r[headers['Ticker']], r[headers['Weight']]
        if not s or s in {'-', 'CASH_USD'} or not isinstance(w, (int, float)) or w <= 0:
            continue
        constituents.append({'symbol': str(s), 'name': str(r[headers['Name']]), 'weight': w / 100,
                             'sector': str(r[headers['Sector']]) if 'Sector' in headers and r[headers['Sector']] not in {None, '-'} else 'Unclassified'})
    return normalize_snapshot({'name': 'SPTM holdings — S&P Composite 1500 proxy', 'as_of': as_of,
                               'source': URL, 'rows': constituents})
